from __future__ import annotations

import re
from html import escape
from pathlib import Path

from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import PageBreak, Paragraph, Preformatted, SimpleDocTemplate, Spacer


ROOT = Path(__file__).resolve().parents[1]
ARTIFACTS = ROOT / "Projeto_Final_Artefatos"
REPORT_SOURCE = ROOT / "docs" / "RELATORIO_TECNICO.md"

NAVY = "12304A"
TEAL = "18C2F2"
CORAL = "F26B5B"
LIGHT = "F3F6F8"
MUTED = "52616B"
TEAM_NAME = "LatinTech"
MEMBERS = "Fábio Castro · Lucas Godois · Pedro Campos · Rogério Sousa · Fernando Gonçalves"
REPOSITORY_URL = "github.com/RogerioSousa123/insurminds-projeto-final-latintech"


def inline_markup(text: str) -> str:
    safe = escape(text)
    safe = re.sub(r"\*\*(.+?)\*\*", r"<b>\1</b>", safe)
    safe = re.sub(r"`(.+?)`", r'<font name="Courier">\1</font>', safe)
    return safe


def generate_report_pdf() -> Path:
    output = ARTIFACTS / "InsurMinds_Relatorio_Tecnico.pdf"
    styles = getSampleStyleSheet()
    styles.add(ParagraphStyle(name="ReportTitle", parent=styles["Title"], alignment=TA_CENTER, textColor=colors.HexColor("#" + NAVY), fontSize=24, leading=29, spaceAfter=10))
    styles.add(ParagraphStyle(name="ReportH1", parent=styles["Heading1"], textColor=colors.HexColor("#" + NAVY), fontSize=17, leading=21, spaceBefore=12, spaceAfter=7))
    styles.add(ParagraphStyle(name="ReportH2", parent=styles["Heading2"], textColor=colors.HexColor("#" + TEAL), fontSize=13, leading=16, spaceBefore=9, spaceAfter=5))
    styles.add(ParagraphStyle(name="ReportH3", parent=styles["Heading3"], textColor=colors.HexColor("#" + NAVY), fontSize=10.5, leading=13, spaceBefore=6, spaceAfter=3))
    styles.add(ParagraphStyle(name="ReportBody", parent=styles["BodyText"], fontSize=9.2, leading=13, spaceAfter=5, textColor=colors.HexColor("#1F2D36")))
    styles.add(ParagraphStyle(name="ReportBullet", parent=styles["BodyText"], fontSize=9, leading=12.5, leftIndent=12, firstLineIndent=-7, spaceAfter=2.5))
    styles.add(ParagraphStyle(name="ReportNote", parent=styles["BodyText"], fontSize=7.8, leading=10, textColor=colors.HexColor("#" + MUTED)))
    code_style = ParagraphStyle(name="Code", fontName="Courier", fontSize=7, leading=9, leftIndent=8, rightIndent=8, backColor=colors.HexColor("#" + LIGHT), borderPadding=6)

    story = []
    paragraph_lines: list[str] = []
    code_lines: list[str] = []
    in_code = False

    def flush_paragraph():
        if paragraph_lines:
            text = " ".join(line.strip() for line in paragraph_lines).strip()
            if text:
                story.append(Paragraph(inline_markup(text), styles["ReportBody"]))
            paragraph_lines.clear()

    for raw in REPORT_SOURCE.read_text(encoding="utf-8").splitlines():
        line = raw.rstrip()
        if line.startswith("```"):
            if in_code:
                story.append(Preformatted("\n".join(code_lines), code_style))
                story.append(Spacer(1, 3 * mm))
                code_lines.clear()
                in_code = False
            else:
                flush_paragraph()
                in_code = True
            continue
        if in_code:
            code_lines.append(line)
            continue
        if line == "---":
            flush_paragraph()
            story.append(PageBreak())
            continue
        if not line.strip():
            flush_paragraph()
            continue
        if line.startswith("### "):
            flush_paragraph()
            story.append(Paragraph(inline_markup(line[4:]), styles["ReportH3"]))
        elif line.startswith("## "):
            flush_paragraph()
            story.append(Paragraph(inline_markup(line[3:]), styles["ReportH1"]))
        elif line.startswith("# "):
            flush_paragraph()
            story.append(Spacer(1, 48 * mm))
            story.append(Paragraph(inline_markup(line[2:]), styles["ReportTitle"]))
        elif re.match(r"^\d+\.\s", line) or line.startswith("- "):
            flush_paragraph()
            content = re.sub(r"^(?:\d+\.|-)\s+", "", line)
            story.append(Paragraph("• " + inline_markup(content), styles["ReportBullet"]))
        elif line.startswith("|"):
            flush_paragraph()
            if not re.match(r"^\|[\s|:-]+\|$", line):
                cells = [cell.strip() for cell in line.strip("|").split("|")]
                story.append(Paragraph(" · ".join(inline_markup(cell) for cell in cells), styles["ReportNote"]))
        else:
            paragraph_lines.append(line)
    flush_paragraph()

    doc = SimpleDocTemplate(
        str(output),
        pagesize=A4,
        rightMargin=22 * mm,
        leftMargin=22 * mm,
        topMargin=18 * mm,
        bottomMargin=18 * mm,
        title="InsurMinds — Relatório Técnico",
        author=TEAM_NAME,
    )

    def footer(canvas, doc_obj):
        canvas.saveState()
        canvas.setStrokeColor(colors.HexColor("#DCE4EA"))
        canvas.line(22 * mm, 14 * mm, A4[0] - 22 * mm, 14 * mm)
        canvas.setFillColor(colors.HexColor("#" + MUTED))
        canvas.setFont("Helvetica", 7)
        canvas.drawString(22 * mm, 9 * mm, "InsurMinds · Projeto Final I2A2 · 2026")
        canvas.drawRightString(A4[0] - 22 * mm, 9 * mm, f"Página {doc_obj.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output


def generate_pitch_deck() -> Path:
    from pptx import Presentation
    from pptx.dml.color import RGBColor
    from pptx.enum.shapes import MSO_SHAPE
    from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
    from pptx.util import Inches, Pt

    output = ARTIFACTS / "InsurMinds_Projeto_Final.pptx"
    presentation = Presentation()
    presentation.slide_width = Inches(13.333)
    presentation.slide_height = Inches(7.5)

    palette = {
        "navy": RGBColor.from_string(NAVY),
        "teal": RGBColor.from_string(TEAL),
        "coral": RGBColor.from_string(CORAL),
        "light": RGBColor.from_string(LIGHT),
        "muted": RGBColor.from_string(MUTED),
        "white": RGBColor(255, 255, 255),
        "dark": RGBColor.from_string("1F2D36"),
        "line": RGBColor.from_string("DCE4EA"),
    }

    def add_text(slide, text, x, y, w, h, size=18, color="dark", bold=False, align=PP_ALIGN.LEFT, font="Aptos", valign=MSO_ANCHOR.TOP):
        box = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        frame = box.text_frame
        frame.clear()
        frame.word_wrap = True
        frame.vertical_anchor = valign
        paragraph = frame.paragraphs[0]
        paragraph.text = text
        paragraph.alignment = align
        paragraph.font.name = font
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = palette[color]
        return box

    def base_slide(title, kicker, number, dark=False):
        slide = presentation.slides.add_slide(presentation.slide_layouts[6])
        background = slide.background.fill
        background.solid()
        background.fore_color.rgb = palette["navy" if dark else "white"]
        add_text(slide, kicker.upper(), 0.7, 0.36, 6, 0.3, 10, "teal" if not dark else "white", True)
        add_text(slide, title, 0.7, 0.74, 11.8, 0.72, 27, "navy" if not dark else "white", True)
        line = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.7), Inches(1.52), Inches(1.1), Inches(0.06))
        line.fill.solid(); line.fill.fore_color.rgb = palette["coral"]; line.line.fill.background()
        add_text(slide, f"{number:02d}", 12.25, 0.35, 0.4, 0.3, 9, "muted" if not dark else "white", True, PP_ALIGN.RIGHT)
        if not dark:
            slide.shapes.add_picture(str(ROOT / "assets" / "logo-latin-re.png"), Inches(11.05), Inches(0.13), width=Inches(0.95))
        return slide

    def card(slide, x, y, w, h, heading, body, accent="teal", number=None):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        shape.fill.solid(); shape.fill.fore_color.rgb = palette["light"]
        shape.line.color.rgb = palette["line"]
        stripe = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x), Inches(y), Inches(0.08), Inches(h))
        stripe.fill.solid(); stripe.fill.fore_color.rgb = palette[accent]; stripe.line.fill.background()
        if number is not None:
            add_text(slide, str(number), x + 0.25, y + 0.18, 0.35, 0.35, 17, accent, True)
            heading_x = x + 0.68
        else:
            heading_x = x + 0.28
        add_text(slide, heading, heading_x, y + 0.18, w - (heading_x - x) - 0.2, 0.36, 15, "navy", True)
        add_text(slide, body, x + 0.28, y + 0.72, w - 0.52, h - 0.9, 10.5, "muted")

    slide = base_slide("Apólices complexas. Comparações explicáveis.", "InsurMinds · Projeto Final", 1, dark=True)
    add_text(slide, "Plataforma inteligente para análise e comparação de seguros D&O", 0.7, 2.05, 7.3, 1.15, 29, "white", True)
    add_text(slide, "Extração estruturada · Evidência por página · Comparação assistida por IA", 0.7, 3.45, 7.6, 0.65, 17, "teal")
    shield = slide.shapes.add_shape(MSO_SHAPE.HEXAGON, Inches(9.5), Inches(2.15), Inches(2.35), Inches(2.55))
    shield.fill.solid(); shield.fill.fore_color.rgb = palette["teal"]; shield.line.fill.background()
    add_text(slide, "D&O", 9.82, 2.82, 1.7, 0.7, 30, "white", True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, f"{TEAM_NAME.upper()} · {MEMBERS}", 0.7, 6.68, 11.8, 0.35, 9, "white", True)

    slide = base_slide("Comparar apólices ainda é um trabalho manual", "Problema", 2)
    card(slide, 0.7, 2.0, 3.85, 3.5, "Documentos extensos", "Cláusulas, limites e exclusões aparecem em páginas diferentes e com terminologia pouco padronizada.", "coral", 1)
    card(slide, 4.75, 2.0, 3.85, 3.5, "Risco de omissão", "Uma diferença relevante pode ficar escondida em condições particulares, exclusões ou endossos.", "coral", 2)
    card(slide, 8.8, 2.0, 3.85, 3.5, "Resultado difícil de auditar", "Um resumo sem página e trecho não permite que o especialista confira rapidamente a conclusão.", "coral", 3)
    add_text(slide, "O objetivo não é retirar o especialista — é reduzir leitura repetitiva e aumentar rastreabilidade.", 1.3, 6.05, 10.7, 0.55, 17, "navy", True, PP_ALIGN.CENTER)

    slide = base_slide("Da leitura dispersa à comparação verificável", "Proposta de valor", 3)
    steps = [
        ("1", "Enviar", "PDF ou imagem"),
        ("2", "Extrair", "texto + OCR"),
        ("3", "Estruturar", "26 critérios D&O"),
        ("4", "Validar", "página + trecho"),
        ("5", "Comparar", "matriz e atenção"),
        ("6", "Consultar", "copiloto fundamentado"),
    ]
    for index, (num, heading, body) in enumerate(steps):
        x = 0.65 + index * 2.08
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(x + 0.55), Inches(2.0), Inches(0.7), Inches(0.7))
        circle.fill.solid(); circle.fill.fore_color.rgb = palette["teal"]; circle.line.fill.background()
        add_text(slide, num, x + 0.55, 2.05, 0.7, 0.55, 17, "white", True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
        add_text(slide, heading, x, 2.95, 1.8, 0.4, 15, "navy", True, PP_ALIGN.CENTER)
        add_text(slide, body, x, 3.42, 1.8, 0.75, 10.5, "muted", False, PP_ALIGN.CENTER)
        if index < len(steps) - 1:
            connector = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x + 1.32), Inches(2.33), Inches(0.83), Inches(0.04))
            connector.fill.solid(); connector.fill.fore_color.rgb = palette["line"]; connector.line.fill.background()
    add_text(slide, "Diferencial: um dado sem evidência verificável não entra na base.", 1.4, 5.55, 10.5, 0.7, 21, "coral", True, PP_ALIGN.CENTER)

    slide = base_slide("Arquitetura híbrida: IA onde agrega valor", "Arquitetura", 4)
    nodes = [
        (0.65, "Documento", "PDF / imagem"),
        (2.65, "Leitura", "PyMuPDF + OCR"),
        (4.65, "Extração IA", "API conversacional"),
        (6.65, "Validação", "Pydantic + evidência"),
        (8.65, "Dados", "SQLite local"),
        (10.65, "Experiência", "Streamlit + relatório"),
    ]
    for index, (x, heading, body) in enumerate(nodes):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(2.15), Inches(1.65), Inches(1.5))
        shape.fill.solid(); shape.fill.fore_color.rgb = palette["navy" if index in {2, 5} else "light"]
        shape.line.color.rgb = palette["teal" if index in {2, 5} else "line"]
        add_text(slide, heading, x + 0.1, 2.46, 1.45, 0.35, 13, "white" if index in {2, 5} else "navy", True, PP_ALIGN.CENTER)
        add_text(slide, body, x + 0.1, 2.9, 1.45, 0.45, 9, "white" if index in {2, 5} else "muted", False, PP_ALIGN.CENTER)
        if index < len(nodes) - 1:
            add_text(slide, "→", x + 1.65, 2.62, 0.35, 0.4, 16, "teal", True, PP_ALIGN.CENTER)
    card(slide, 1.0, 4.45, 3.45, 1.35, "Código determinístico", "Datas, valores, estados, cache e diferenças reproduzíveis.", "teal")
    card(slide, 4.95, 4.45, 3.45, 1.35, "IA generativa", "Interpretação de cláusulas, síntese e perguntas em linguagem natural.", "coral")
    card(slide, 8.9, 4.45, 3.45, 1.35, "Humano no controle", "Evidências, ambiguidades e ressalvas ficam visíveis para conferência.", "teal")

    slide = base_slide("Agentes pequenos, especializados e testáveis", "Orquestração", 5)
    agents = [
        ("Leitor", "Valida arquivo e preserva texto por página."),
        ("Extrator", "Mapeia cláusulas à taxonomia fechada."),
        ("Auditor", "Rejeita página ou trecho não confirmável."),
        ("Comparador", "Confronta valores e interpreta diferenças."),
        ("Relator", "Gera PDF, CSV e JSON auditáveis."),
        ("Copiloto", "Responde somente com contexto recuperado."),
    ]
    for index, (heading, body) in enumerate(agents):
        row, col = divmod(index, 3)
        card(slide, 0.75 + col * 4.18, 1.95 + row * 2.15, 3.82, 1.75, heading, body, "teal" if index % 2 == 0 else "coral", index + 1)

    slide = base_slide("Confiabilidade construída fora do prompt", "Controles", 6)
    controls = [
        ("Taxonomia fechada", "Chaves desconhecidas são descartadas."),
        ("Evidência obrigatória", "Página e trecho são verificados no texto original."),
        ("Estados distintos", "Não localizado ≠ não coberto; ambiguidade fica explícita."),
        ("Proteção de credenciais", "Chaves em .env, nunca no banco ou Git."),
        ("Cache por hash", "Evita custo e reprocessamento do mesmo documento."),
        ("Prompt injection", "Documento é tratado como dado não confiável."),
    ]
    for index, (heading, body) in enumerate(controls):
        y = 1.85 + index * 0.78
        dot = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.9), Inches(y + 0.04), Inches(0.35), Inches(0.35))
        dot.fill.solid(); dot.fill.fore_color.rgb = palette["teal" if index % 2 == 0 else "coral"]; dot.line.fill.background()
        add_text(slide, heading, 1.5, y, 3.0, 0.35, 14, "navy", True)
        add_text(slide, body, 4.45, y, 7.7, 0.45, 12, "muted")

    slide = base_slide("MVP funcional e demonstrável", "Resultados", 7)
    metrics = [
        ("26", "critérios D&O estruturados"),
        ("4", "apólices por comparação"),
        ("14/14", "testes automatizados aprovados"),
        ("3", "formatos de exportação"),
    ]
    for index, (number, label) in enumerate(metrics):
        x = 0.7 + index * 3.15
        add_text(slide, number, x, 2.0, 2.7, 0.9, 34, "teal" if index % 2 == 0 else "coral", True, PP_ALIGN.CENTER)
        add_text(slide, label, x, 3.05, 2.7, 0.8, 13, "navy", True, PP_ALIGN.CENTER)
    card(slide, 1.05, 4.65, 3.45, 1.15, "Entrada", "PDF, PNG, JPEG, TIFF e WebP", "teal")
    card(slide, 4.95, 4.65, 3.45, 1.15, "Saída", "Matriz, resumo, evidências e chat", "coral")
    card(slide, 8.85, 4.65, 3.45, 1.15, "Contingência", "Persistência local e modo demo explícito", "teal")

    slide = base_slide("Limitações conhecidas orientam a evolução", "Roadmap", 8)
    card(slide, 0.75, 1.9, 3.8, 3.75, "Hoje", "• OCR depende da qualidade\n• Revisão jurídica é obrigatória\n• Busca lexical no copiloto\n• SQLite e usuário local\n• Taxonomia focada em D&O", "coral")
    card(slide, 4.78, 1.9, 3.8, 3.75, "Próximos passos", "• Revisão humana assistida\n• Métricas em conjunto rotulado\n• Busca híbrida / embeddings\n• Endossos e apólice-base\n• Custos por documento", "teal")
    card(slide, 8.81, 1.9, 3.8, 3.75, "Visão futura", "• Perfis e trilha de auditoria\n• Novos produtos de seguro\n• Implantação privada\n• Integração com corretoras\n• Aprendizado com correções", "teal")

    slide = base_slide("Uma decisão melhor começa por uma evidência melhor", "Encerramento", 9, dark=True)
    add_text(slide, "A InsurMinds transforma apólices complexas em uma comparação estruturada, explicável e verificável — sem retirar o especialista da decisão.", 1.25, 2.0, 10.8, 1.65, 27, "white", True, PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    add_text(slide, "PROBLEMA  →  EVIDÊNCIA  →  COMPARAÇÃO  →  DECISÃO HUMANA", 1.4, 4.25, 10.5, 0.6, 16, "teal", True, PP_ALIGN.CENTER)
    add_text(slide, f"GRUPO {TEAM_NAME.upper()} · {MEMBERS}", 1.0, 6.12, 11.3, 0.34, 9, "white", True, PP_ALIGN.CENTER)
    add_text(slide, REPOSITORY_URL, 1.0, 6.53, 11.3, 0.3, 9, "teal", True, PP_ALIGN.CENTER)

    presentation.save(str(output))
    return output


def main():
    ARTIFACTS.mkdir(parents=True, exist_ok=True)
    report = generate_report_pdf()
    deck = generate_pitch_deck()
    print(f"Relatório: {report}")
    print(f"Pitch deck: {deck}")


if __name__ == "__main__":
    main()
